from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from typing import List, Optional, Dict, Generator
import os
import requests
import json
import re

import pdfplumber
import docx


# ==============================
# CONFIG
# ==============================

app = FastAPI(title="ResumeBuilderAI - Qwen Edition")

OLLAMA_API = os.getenv("OLLAMA_API", "http://127.0.0.1:11434/api/generate")
MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:7b-instruct-q4_K_M")
OLLAMA_TIMEOUT = int(os.getenv("OLLAMA_TIMEOUT", "600"))

# --- Hosted LLM (OpenAI-compatible, e.g. Groq) --------------------------------
# When LLM_API_KEY is set, the service calls a hosted OpenAI-compatible Chat
# Completions API instead of local Ollama. This is fast and needs NO GPU, which
# makes it the right fit for CPU-only hosts like Railway. When the key is empty
# we transparently fall back to Ollama (used by the docker-compose dev stack).
#   Groq:  LLM_BASE_URL=https://api.groq.com/openai/v1  LLM_MODEL=llama-3.3-70b-versatile
LLM_API_KEY = os.getenv("LLM_API_KEY") or os.getenv("GROQ_API_KEY", "")
LLM_BASE_URL = os.getenv("LLM_BASE_URL", "https://api.groq.com/openai/v1").rstrip("/")
# 2026-08-28: Groq retired `llama-3.3-70b-versatile` (API now returns 404
# model_not_found for it). `openai/gpt-oss-120b` is the closest replacement on the
# same free tier and supports response_format=json_object. Override via LLM_MODEL.
LLM_MODEL = os.getenv("LLM_MODEL", "openai/gpt-oss-120b")
LLM_TIMEOUT = int(os.getenv("LLM_TIMEOUT", "120"))
USE_HOSTED_LLM = bool(LLM_API_KEY)


# ==============================
# DATA MODELS
# ==============================

class Experience(BaseModel):
    role: str
    company: str
    location: Optional[str] = None
    startDate: str
    endDate: str
    bullets: List[str] = Field(default_factory=list)


class Project(BaseModel):
    title: str
    link: Optional[str] = None
    bullets: List[str] = Field(default_factory=list)


class Education(BaseModel):
    school: str
    degree: str
    field: str
    location: Optional[str] = None
    endDate: str


class SkillCategory(BaseModel):
    category: str
    skills: List[str]


class ResumeRequest(BaseModel):
    name: str
    email: str
    phone: str
    linkedin: str
    portfolio: str
    summary: str
    experiences: List[Experience]
    projects: List[Project]
    education: List[Education]
    skills: List[SkillCategory]
    job_description: str


class OptimizedResume(BaseModel):
    summary: str
    experiences: List[Experience]
    projects: List[Project]
    education: List[Education]
    skills: List[SkillCategory]
    ats_report: Dict[str, object]


# ==============================
# OLLAMA CALL
# ==============================

def _hosted_headers():
    return {"Authorization": f"Bearer {LLM_API_KEY}", "Content-Type": "application/json"}


def _hosted_complete(prompt: str, *, temperature: float = 0.15, json_mode: bool = False) -> str:
    """Blocking chat-completion against the hosted OpenAI-compatible API (returns text)."""
    payload = {
        "model": LLM_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": temperature,
    }
    if json_mode:
        payload["response_format"] = {"type": "json_object"}
    resp = requests.post(
        f"{LLM_BASE_URL}/chat/completions",
        headers=_hosted_headers(),
        json=payload,
        timeout=LLM_TIMEOUT,
    )
    resp.raise_for_status()
    data = resp.json()
    return ((data.get("choices") or [{}])[0].get("message") or {}).get("content", "") or ""


def _hosted_stream(prompt: str, temperature: float = 0.6) -> Generator[bytes, None, None]:
    """Streaming chat-completion against the hosted OpenAI-compatible API (SSE)."""
    payload = {
        "model": LLM_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": temperature,
        "stream": True,
    }
    with requests.post(
        f"{LLM_BASE_URL}/chat/completions",
        headers=_hosted_headers(),
        json=payload,
        stream=True,
        timeout=LLM_TIMEOUT,
    ) as resp:
        resp.raise_for_status()
        for line in resp.iter_lines(decode_unicode=True):
            if not line:
                continue
            if line.startswith("data:"):
                data = line[5:].strip()
                if data == "[DONE]":
                    break
                try:
                    obj = json.loads(data)
                except Exception:
                    continue
                delta = ((obj.get("choices") or [{}])[0].get("delta") or {})
                chunk = delta.get("content")
                if chunk:
                    yield chunk.encode("utf-8")


def call_ollama(prompt: str):
    # Hosted API (Groq etc.) when configured; otherwise local Ollama.
    if USE_HOSTED_LLM:
        return _hosted_complete(prompt, temperature=0.15, json_mode=True)

    payload = {
        "model": MODEL,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {
            "temperature": 0.15,
            "top_p": 0.9,
            "num_ctx": 4096
        }
    }

    response = requests.post(OLLAMA_API, json=payload, timeout=OLLAMA_TIMEOUT)
    response.raise_for_status()

    return response.json().get("response", "")


def safe_json_parse(text: str):

    try:
        return json.loads(text)

    except:

        match = re.search(r"\{.*\}", text, re.DOTALL)

        if not match:
            raise ValueError("Model did not return valid JSON.")

        return json.loads(match.group(0))


# ==============================
# RESUME TEXT EXTRACTION
# ==============================

def extract_resume_text(file: UploadFile):

    filename = file.filename.lower()

    if filename.endswith(".pdf"):

        with pdfplumber.open(file.file) as pdf:

            text = ""

            for page in pdf.pages:
                text += page.extract_text() or ""

        return text


    elif filename.endswith(".docx"):

        doc = docx.Document(file.file)

        text = "\n".join([p.text for p in doc.paragraphs])

        return text

    else:
        raise HTTPException(
            status_code=400,
            detail="Only PDF or DOCX resumes supported"
        )


# ==============================
# RESUME SECTION SPLITTER
# ==============================

def split_resume_sections(text: str):

    sections = {
        "summary": "",
        "experience": "",
        "education": "",
        "projects": "",
        "skills": ""
    }

    patterns = {

        "summary": r"(summary|professional summary|profile|about me|career summary|career overview|objective|personal statement)",

        "experience": r"(experience|employment|work history|professional experience|career history|internships)",

        "education": r"(education|academic|academic background|qualifications)",

        "projects": r"(projects|portfolio|personal projects|research)",

        "skills": r"(skills|technical skills|core skills|competencies|technologies|tech stack)"
    }

    lines = text.split("\n")

    current = "summary"   # safe default

    for line in lines:

        lower = line.lower().strip()

        for key, pattern in patterns.items():
            if re.search(pattern, lower):
                current = key
                continue

        if current in sections:
            sections[current] += line + "\n"

    return sections


def extract_contact_info(text: str):

    email_match = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)

    phone_match = re.search(r"\+?\d[\d\s\-]{8,15}", text)

    email = email_match.group(0) if email_match else ""
    phone = phone_match.group(0) if phone_match else ""

    return email, phone


# ==============================
# RESUME PARSER
# ==============================

def parse_resume_to_schema(resume_text: str):

    sections = split_resume_sections(resume_text)
    email, phone = extract_contact_info(resume_text)

    prompt = f"""
You are a resume parser.

Convert the resume sections into the EXACT JSON schema below.

If information missing:
- use empty string ""
- or empty lists []

Do NOT invent data.

Return STRICT JSON ONLY.

SCHEMA:

{{
  "name": "",
  "email": "",
  "phone": "",
  "linkedin": "",
  "portfolio": "",
  "summary": "",
  "experiences": [
    {{
      "role": "",
      "company": "",
      "location": "",
      "startDate": "",
      "endDate": "",
      "bullets": []
    }}
  ],
  "projects": [
    {{
      "title": "",
      "link": "",
      "bullets": []
    }}
  ],
  "education": [
    {{
      "school": "",
      "degree": "",
      "field": "",
      "location": "",
      "endDate": ""
    }}
  ],
  "skills": [
    {{
      "category": "",
      "skills": []
    }}
  ],
  "job_description": ""
}}

RESUME SECTIONS:
Detected email: {email}
Detected phone: {phone}

If these values exist, use them in the JSON.
SUMMARY:
{sections["summary"]}

EXPERIENCE:
{sections["experience"]}

EDUCATION:
{sections["education"]}

PROJECTS:
{sections["projects"]}

SKILLS:
{sections["skills"]}
"""

    result = call_ollama(prompt)

    parsed = safe_json_parse(result)

    return parsed


# ==============================
# PASS 1: EXTRACT JD SIGNALS
# ==============================

def extract_jd_signals(job_description: str):

    prompt = f"""
You are an ATS job analyzer.

Return STRICT JSON only:

{{
  "must_have_keywords": ["..."],
  "nice_to_have_keywords": ["..."],
  "tools": ["..."],
  "responsibilities": ["..."],
  "soft_skills": ["..."]
}}

Rules:
- Extract 10-20 must-have ATS keywords.
- Keep terms short (e.g., "SQL", "Power BI", "ETL").
- No commentary. JSON only.

JOB DESCRIPTION:
{job_description}
"""

    result = call_ollama(prompt)

    return safe_json_parse(result)


# ==============================
# PASS 2: OPTIMIZE RESUME
# ==============================

def optimize_resume(req: ResumeRequest, jd_data: dict):

    prompt = f"""
You are ResumeBuilderAI.

Rewrite the resume to maximize ATS match using the extracted job signals.

Return STRICT JSON only:

{{
  "summary": "...",
  "experiences": [...],
  "projects": [...],
  "education": [...],
  "skills": [...],
  "ats_report": {{
      "coverage_percent": 0,
      "keywords_covered": ["..."],
      "keywords_missing": ["..."],
      "notes": ["..."]
  }}
}}

Rules:
- Do not invent employers or degrees.
- Improve impact with quantified achievements.
- Naturally integrate must-have keywords.
- 3–6 bullets per experience.
- Professional, ATS-safe language.

JOB SIGNALS:
{json.dumps(jd_data, indent=2)}

CANDIDATE RESUME:
{req.model_dump_json(indent=2)}
"""

    result = call_ollama(prompt)

    parsed = safe_json_parse(result)

    return OptimizedResume(**parsed)


# ==============================
# ATS SCORING HELPERS
# ==============================

def normalize_text(text: str):
    return re.sub(r"\s+", " ", text.lower()).strip()


def keyword_match_score(resume_text: str, keywords: list):
    """
    Calculates keyword coverage percentage.

    Returns:
    - score
    - covered keywords
    - missing keywords
    """

    normalized_resume = normalize_text(resume_text)

    keywords = list(dict.fromkeys([
        k.strip().lower()
        for k in keywords
        if isinstance(k, str) and k.strip()
    ]))

    covered = []
    missing = []

    for keyword in keywords:
        keyword_norm = normalize_text(keyword)

        if keyword_norm in normalized_resume:
            covered.append(keyword)
        else:
            missing.append(keyword)

    score = round((len(covered) / max(len(keywords), 1)) * 100, 2)

    return score, covered, missing


def responsibility_alignment_score(resume_text: str, responsibilities: list):
    """
    Scores whether resume content reflects job responsibilities,
    not just isolated keywords.
    """

    normalized_resume = normalize_text(resume_text)

    responsibilities = [
        r.strip().lower()
        for r in responsibilities
        if isinstance(r, str) and r.strip()
    ]

    if not responsibilities:
        return 60

    matched = 0

    stop_words = {
        "with", "from", "that", "this", "will", "have",
        "your", "their", "using", "work", "team", "role",
        "able", "into", "over", "under", "within", "across",
        "while", "also", "such", "including", "based"
    }

    for responsibility in responsibilities:

        words = [
            w for w in re.findall(r"\b[a-zA-Z]{4,}\b", responsibility)
            if w not in stop_words
        ]

        if not words:
            continue

        overlap = sum(1 for w in words if w in normalized_resume)

        if overlap / max(len(words), 1) >= 0.35:
            matched += 1

    return round((matched / max(len(responsibilities), 1)) * 100, 2)


def quantified_achievement_score(resume: OptimizedResume):
    """
    Rewards bullets with measurable achievements, metrics,
    or impact-focused wording.
    """

    all_bullets = []

    for exp in resume.experiences:
        all_bullets.extend(exp.bullets)

    for project in resume.projects:
        all_bullets.extend(project.bullets)

    if not all_bullets:
        return 40

    quantified_count = 0

    impact_terms = [
        "increased",
        "reduced",
        "improved",
        "optimized",
        "decreased",
        "saved",
        "automated",
        "accelerated",
        "enhanced",
        "delivered",
        "achieved",
        "boosted",
        "streamlined",
        "cut",
        "grew",
        "resolved"
    ]

    for bullet in all_bullets:

        bullet_lower = bullet.lower()

        has_number = bool(re.search(r"\d+|%|percent", bullet_lower))
        has_impact = any(term in bullet_lower for term in impact_terms)

        if has_number or has_impact:
            quantified_count += 1

    score = round((quantified_count / max(len(all_bullets), 1)) * 100, 2)

    return min(score, 100)


def structure_score(resume: OptimizedResume):
    """
    Checks resume completeness.
    """

    score = 0

    if resume.summary and len(resume.summary.split()) >= 25:
        score += 20

    if resume.experiences:
        score += 25

    if resume.projects:
        score += 20

    if resume.education:
        score += 15

    if resume.skills:
        score += 20

    return score


def keyword_stuffing_penalty(resume: OptimizedResume, jd_data: dict):
    """
    Penalizes unnatural keyword repetition.
    """

    resume_text = normalize_text(json.dumps(resume.model_dump()))

    all_keywords = (
        jd_data.get("must_have_keywords", []) +
        jd_data.get("nice_to_have_keywords", []) +
        jd_data.get("tools", [])
    )

    penalty = 0

    for keyword in all_keywords:

        keyword_norm = normalize_text(keyword)

        if not keyword_norm:
            continue

        count = resume_text.count(keyword_norm)

        if count >= 5:
            penalty += 2

    return min(penalty, 5)


# ==============================
# PASS 3: COVERAGE CHECK
# ==============================

def compute_coverage(resume: OptimizedResume, jd_data: dict):

    must_have = jd_data.get("must_have_keywords", [])

    resume_text = json.dumps(resume.model_dump())

    keyword_score, covered, missing = keyword_match_score(
        resume_text,
        must_have
    )

    return keyword_score, covered, missing


def improve_resume_if_needed(resume: OptimizedResume, jd_data: dict):

    coverage, covered, missing = compute_coverage(resume, jd_data)

    if coverage >= 85:
        return resume

    prompt = f"""
The resume below is missing important ATS keywords:

{missing}

Improve the resume by:
- Naturally integrating missing keywords
- Strengthening quantified achievements
- Keeping same structure
- Not inventing companies or degrees

Return STRICT JSON in same format as before.

CURRENT RESUME:
{resume.model_dump_json(indent=2)}

JOB SIGNALS:
{json.dumps(jd_data, indent=2)}
"""

    improved = safe_json_parse(call_ollama(prompt))

    return OptimizedResume(**improved)


# ==============================
# ATS KEYWORD + QUALITY SCORING
# ==============================

def compute_ats_score(resume: OptimizedResume, jd_data: dict):

    resume_text = json.dumps(resume.model_dump())

    must_score, must_covered, must_missing = keyword_match_score(
        resume_text,
        jd_data.get("must_have_keywords", [])
    )

    nice_tools_score, nice_tools_covered, nice_tools_missing = keyword_match_score(
        resume_text,
        jd_data.get("nice_to_have_keywords", []) + jd_data.get("tools", [])
    )

    responsibility_score = responsibility_alignment_score(
        resume_text,
        jd_data.get("responsibilities", [])
    )

    quantified_score = quantified_achievement_score(resume)

    structural_score = structure_score(resume)

    stuffing_penalty = keyword_stuffing_penalty(resume, jd_data)

    raw_score = (
        (must_score * 0.35) +
        (nice_tools_score * 0.15) +
        (responsibility_score * 0.20) +
        (quantified_score * 0.15) +
        (structural_score * 0.10) -
        stuffing_penalty
    )

    final_score = round(raw_score, 2)

    # Realistic cap so the score behaves more like online ATS checkers
    final_score = max(0, min(final_score, 89))

    return {
        "final_ats_score": final_score,

        # Same keys as before, so frontend remains compatible
        "keywords_found": must_covered + nice_tools_covered,
        "keywords_missing": must_missing + nice_tools_missing,

        # Extra details; frontend can ignore this if not used
        "score_breakdown": {
            "must_have_keyword_score": must_score,
            "nice_to_have_tools_score": nice_tools_score,
            "responsibility_alignment_score": responsibility_score,
            "quantified_achievement_score": quantified_score,
            "structure_score": structural_score,
            "keyword_stuffing_penalty": stuffing_penalty
        }
    }


# ==============================
# IMPROVEMENT LOOP
# ==============================

def optimize_until_threshold(resume: OptimizedResume, jd_data: dict, threshold=78):

    max_iterations = 3

    for i in range(max_iterations):

        ats = compute_ats_score(resume, jd_data)

        if ats["final_ats_score"] >= threshold:
            return resume, ats, i + 1

        missing = ats["keywords_missing"]
        breakdown = ats.get("score_breakdown", {})

        prompt = f"""
The resume below has an ATS score of {ats["final_ats_score"]}.

It must reach at least {threshold}.

Current score breakdown:
{json.dumps(breakdown, indent=2)}

Missing keywords:
{missing}

Improve the resume by:
- Integrating missing keywords naturally
- Improving responsibility alignment with the job description
- Strengthening quantified achievements
- Making bullets more specific and impact-focused
- Maintaining the same JSON structure
- Not inventing companies, degrees, job titles, dates, or fake experience

Return STRICT JSON in the same schema.

CURRENT RESUME:
{resume.model_dump_json(indent=2)}

JOB SIGNALS:
{json.dumps(jd_data, indent=2)}
"""

        improved = safe_json_parse(call_ollama(prompt))

        resume = OptimizedResume(**improved)

    ats = compute_ats_score(resume, jd_data)

    return resume, ats, max_iterations


# ==============================
# RESUME PARSER ENDPOINT
# ==============================

@app.post("/parse_resume")
async def parse_resume(file: UploadFile = File(...)):

    try:

        resume_text = extract_resume_text(file)

        parsed_resume = parse_resume_to_schema(resume_text)

        return parsed_resume

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# ==============================
# GENERATOR ENDPOINT
# ==============================

@app.post("/generate_resume")
def generate_resume(req: ResumeRequest):

    try:

        jd_data = extract_jd_signals(req.job_description)

        optimized_resume = optimize_resume(req, jd_data)

        optimized_resume = improve_resume_if_needed(
            optimized_resume,
            jd_data
        )

        optimized_resume, ats_score, iterations = optimize_until_threshold(
            optimized_resume,
            jd_data
        )

        return {

            "candidate_info": {
                "name": req.name,
                "email": req.email,
                "phone": req.phone,
                "linkedin": req.linkedin,
                "portfolio": req.portfolio
            },

            "resume": optimized_resume.model_dump(),

            "jd_analysis": jd_data,

            "ats_final_result": {
                "final_ats_score": ats_score["final_ats_score"],
                "iterations_needed": iterations,
                "keywords_found": ats_score["keywords_found"],
                "keywords_missing": ats_score["keywords_missing"],
                "score_breakdown": ats_score.get("score_breakdown", {})
            }

        }

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# ==============================
# COVER LETTER GENERATION
# ==============================

class CoverLetterRequest(BaseModel):
    # Either send a structured resume (preferred) or a plain-text resume.
    resume: Optional[ResumeRequest] = None
    resume_text: Optional[str] = None
    job_description: str
    tone: Optional[str] = "professional"  # professional | enthusiastic | concise | warm
    company: Optional[str] = None
    role: Optional[str] = None


def _resume_to_plaintext(r: ResumeRequest) -> str:
    lines: List[str] = []
    if r.name:
        lines.append(r.name)
    contact_bits = [b for b in [r.email, r.phone, r.linkedin, r.portfolio] if b]
    if contact_bits:
        lines.append(" | ".join(contact_bits))
    if r.summary:
        lines.append(f"\nSUMMARY\n{r.summary}")
    if r.experiences:
        lines.append("\nEXPERIENCE")
        for e in r.experiences:
            header = " — ".join([p for p in [e.role, e.company] if p])
            dates = " – ".join([p for p in [e.startDate, e.endDate] if p])
            lines.append(f"{header} ({dates})" if dates else header)
            for b in e.bullets:
                lines.append(f"- {b}")
    if r.projects:
        lines.append("\nPROJECTS")
        for p in r.projects:
            lines.append(p.title)
            for b in p.bullets:
                lines.append(f"- {b}")
    if r.education:
        lines.append("\nEDUCATION")
        for ed in r.education:
            lines.append(" — ".join([p for p in [ed.degree, ed.field, ed.school] if p]))
    if r.skills:
        lines.append("\nSKILLS")
        for cat in r.skills:
            lines.append(f"{cat.category}: {', '.join(cat.skills)}")
    return "\n".join(lines).strip()


_TONE_HINTS = {
    "professional": "polished, confident, neutral",
    "enthusiastic": "warm, energetic, genuinely excited",
    "concise": "tight, direct, no fluff — under 220 words",
    "warm": "personable, conversational, sincere",
}


def _build_cover_letter_prompt(
    resume_text: str,
    job_description: str,
    tone: str,
    company: Optional[str],
    role: Optional[str],
) -> str:
    tone_hint = _TONE_HINTS.get(tone, _TONE_HINTS["professional"])
    target = ""
    if role and company:
        target = f"\nROLE: {role} at {company}"
    elif role:
        target = f"\nROLE: {role}"
    elif company:
        target = f"\nCOMPANY: {company}"

    return f"""You are an expert career writer. Write a cover letter in {tone_hint} tone.

Rules:
- 3 short paragraphs (opening, fit, close).
- Pull SPECIFIC achievements from the resume that match the job description.
- Do not invent jobs, employers, dates, or skills the candidate doesn't have.
- Do not list bullet points. Write flowing prose.
- Do not include placeholders like [Your Name] or [Date] — only the letter body.
- Do not repeat the resume verbatim.
- Start with "Dear Hiring Manager," and end with "Sincerely," followed by the candidate's name on a new line.
- CRITICAL: Only reference company names, team locations, technologies, or job specifics that are explicitly stated in the JOB DESCRIPTION. Do NOT copy candidate preferences or aspirations from the resume summary (e.g. "open to remote roles with Australian teams") into the letter as if they are facts about this specific job.
{target}

JOB DESCRIPTION:
{job_description}

CANDIDATE RESUME:
{resume_text}

Write the cover letter now:"""


def stream_ollama(prompt: str, temperature: float = 0.6) -> Generator[bytes, None, None]:
    # Hosted API (Groq etc.) when configured; otherwise local Ollama.
    if USE_HOSTED_LLM:
        yield from _hosted_stream(prompt, temperature)
        return

    payload = {
        "model": MODEL,
        "prompt": prompt,
        "stream": True,
        "options": {
            "temperature": temperature,
            "top_p": 0.9,
            "num_ctx": 4096,
        },
    }
    with requests.post(OLLAMA_API, json=payload, stream=True, timeout=OLLAMA_TIMEOUT) as resp:
        resp.raise_for_status()
        for line in resp.iter_lines():
            if not line:
                continue
            try:
                obj = json.loads(line.decode("utf-8"))
            except Exception:
                continue
            chunk = obj.get("response", "")
            if chunk:
                yield chunk.encode("utf-8")
            if obj.get("done"):
                break


@app.post("/generate_cover_letter")
def generate_cover_letter(req: CoverLetterRequest):
    if not (req.job_description or "").strip():
        raise HTTPException(status_code=400, detail="job_description is required")

    if req.resume is not None:
        resume_text = _resume_to_plaintext(req.resume)
    elif req.resume_text and req.resume_text.strip():
        resume_text = req.resume_text.strip()
    else:
        raise HTTPException(status_code=400, detail="Provide either `resume` or `resume_text`")

    prompt = _build_cover_letter_prompt(
        resume_text=resume_text,
        job_description=req.job_description.strip(),
        tone=(req.tone or "professional").lower(),
        company=req.company,
        role=req.role,
    )

    return StreamingResponse(
        stream_ollama(prompt, temperature=0.6),
        media_type="text/plain; charset=utf-8",
    )


# ==============================
# Q&A ANSWERS GENERATION
# ==============================

class QAAnswersRequest(BaseModel):
    resume: Optional[ResumeRequest] = None
    resume_text: Optional[str] = None
    job_description: str
    tone: Optional[str] = "professional"
    company: Optional[str] = None
    role: Optional[str] = None
    interview_type: Optional[str] = "behavioral"  # screening | behavioral | technical | manager
    focus: Optional[str] = None
    question_count: Optional[int] = 6
    questions: Optional[List[str]] = None


_INTERVIEW_TYPE_HINTS = {
    "screening": "recruiter phone-screen questions (motivation, background basics, logistics, salary)",
    "behavioral": "behavioral questions answered in STAR form (Situation, Task, Action, Result)",
    "technical": "role-specific technical questions that probe depth and trade-offs",
    "manager": "hiring-manager questions about impact, ownership, collaboration, and fit",
}


def _resume_or_text(resume: Optional[ResumeRequest], resume_text: Optional[str]) -> str:
    if resume is not None:
        return _resume_to_plaintext(resume)
    if resume_text and resume_text.strip():
        return resume_text.strip()
    return ""


def _target_line(role: Optional[str], company: Optional[str]) -> str:
    if role and company:
        return f"\nROLE: {role} at {company}"
    if role:
        return f"\nROLE: {role}"
    if company:
        return f"\nCOMPANY: {company}"
    return ""


def _build_qa_prompt(
    resume_text: str,
    job_description: str,
    tone: str,
    company: Optional[str],
    role: Optional[str],
    interview_type: str,
    focus: Optional[str],
    question_count: int,
    questions: Optional[List[str]],
) -> str:
    tone_hint = _TONE_HINTS.get(tone, _TONE_HINTS["professional"])
    type_hint = _INTERVIEW_TYPE_HINTS.get(interview_type, _INTERVIEW_TYPE_HINTS["behavioral"])
    target = _target_line(role, company)
    focus_line = f"\nFOCUS AREAS: {focus}" if focus else ""

    if questions:
        numbered = "\n".join(f"{i + 1}. {q}" for i, q in enumerate(questions))
        task = (
            "Answer EACH of the interview questions below. Write the answer in the FIRST PERSON "
            "as the candidate, grounded in the candidate's real resume and tailored to the job description."
        )
        questions_block = f"\nQUESTIONS TO ANSWER:\n{numbered}\n"
    else:
        task = (
            f"Generate {question_count} likely {type_hint} for this role, and for EACH one write a strong, "
            "specific answer in the FIRST PERSON as the candidate, grounded in the resume and tailored to the job."
        )
        questions_block = ""

    return f"""You are an expert interview coach helping a candidate prepare. {task}

Tone: {tone_hint}.
Interview type: {type_hint}.{target}{focus_line}

Rules:
- Use ONLY real experience from the candidate's resume. Do NOT invent employers, titles, dates, or skills.
- Be concrete: reference specific projects, technologies, and measurable results from the resume.
- Keep each answer roughly 90–160 words. Natural spoken prose, not bullet lists.
- Output PLAIN TEXT only. For each item use exactly this format:

Q: <the question>
A: <the answer>

(blank line between items). No markdown, no numbering, no preamble, no closing remarks.
{questions_block}
JOB DESCRIPTION:
{job_description}

CANDIDATE RESUME:
{resume_text or "(no resume provided — answer generically but plausibly for this role)"}

Begin now:"""


@app.post("/generate_qa_answers")
def generate_qa_answers(req: QAAnswersRequest):
    if not (req.job_description or "").strip():
        raise HTTPException(status_code=400, detail="job_description is required")

    prompt = _build_qa_prompt(
        resume_text=_resume_or_text(req.resume, req.resume_text),
        job_description=req.job_description.strip(),
        tone=(req.tone or "professional").lower(),
        company=req.company,
        role=req.role,
        interview_type=(req.interview_type or "behavioral").lower(),
        focus=req.focus,
        question_count=max(1, min(int(req.question_count or 6), 15)),
        questions=[q for q in (req.questions or []) if q and q.strip()] or None,
    )

    return StreamingResponse(
        stream_ollama(prompt, temperature=0.5),
        media_type="text/plain; charset=utf-8",
    )


# ==============================
# HR EMAIL DRAFTS GENERATION
# ==============================

class HREmailRequest(BaseModel):
    resume: Optional[ResumeRequest] = None
    resume_text: Optional[str] = None
    job_description: Optional[str] = None
    tone: Optional[str] = "professional"
    company: Optional[str] = None
    role: Optional[str] = None
    email_type: Optional[str] = "application"
    recipient_name: Optional[str] = None
    job_link: Optional[str] = None
    date_applied: Optional[str] = None
    availability: Optional[str] = None
    extra_context: Optional[str] = None
    drafts: Optional[int] = 2


_EMAIL_TYPE_HINTS = {
    "application": "a job application email introducing the candidate and expressing interest",
    "follow_up": "a polite follow-up email checking on application status",
    "thank_you": "a thank-you email after an interview",
    "scheduling": "an email proposing/confirming interview scheduling and availability",
    "referral_request": "an email requesting a referral or introduction",
    "offer_clarification": "an email asking clarifying questions about a job offer",
    "negotiation": "a professional salary/offer negotiation email",
}


def _build_hr_email_prompt(
    resume_text: str,
    job_description: Optional[str],
    tone: str,
    company: Optional[str],
    role: Optional[str],
    email_type: str,
    recipient_name: Optional[str],
    job_link: Optional[str],
    date_applied: Optional[str],
    availability: Optional[str],
    extra_context: Optional[str],
    drafts: int,
) -> str:
    tone_hint = _TONE_HINTS.get(tone, _TONE_HINTS["professional"])
    type_hint = _EMAIL_TYPE_HINTS.get(email_type, _EMAIL_TYPE_HINTS["application"])
    target = _target_line(role, company)

    context_bits = []
    if recipient_name:
        context_bits.append(f"RECIPIENT: {recipient_name}")
    if job_link:
        context_bits.append(f"JOB LINK: {job_link}")
    if date_applied:
        context_bits.append(f"DATE APPLIED: {date_applied}")
    if availability:
        context_bits.append(f"AVAILABILITY: {availability}")
    if extra_context:
        context_bits.append(f"EXTRA CONTEXT: {extra_context}")
    context_block = ("\n" + "\n".join(context_bits)) if context_bits else ""

    plural = "drafts" if drafts > 1 else "draft"
    greeting = f"Dear {recipient_name}," if recipient_name else "Dear Hiring Manager,"
    # Precompute the optional job-description block: f-string expression parts can't
    # contain backslashes on Python < 3.12, so build any newline-containing pieces here.
    jd_block = f"JOB DESCRIPTION:\n{job_description}\n\n" if (job_description or "").strip() else ""
    resume_block = resume_text or "(no resume provided — write plausibly using the context above)"

    return f"""You are an expert career communication writer. Write {drafts} distinct {plural} of {type_hint}.

Tone: {tone_hint}.{target}{context_block}

Rules:
- Each draft must be self-contained and ready to send.
- Use ONLY real details from the candidate's resume; do NOT invent facts.
- Keep each email concise (under ~160 words) and specific to the role/company.
- Start the body with "{greeting}" and end with "Best regards," followed by the candidate's name.
- Output PLAIN TEXT only. For each draft use exactly this format:

--- Draft N ---
Subject: <subject line>
<email body>

No markdown, no commentary before or after the drafts.

{jd_block}CANDIDATE RESUME:
{resume_block}

Write the {plural} now:"""


@app.post("/generate_hr_email")
def generate_hr_email(req: HREmailRequest):
    prompt = _build_hr_email_prompt(
        resume_text=_resume_or_text(req.resume, req.resume_text),
        job_description=req.job_description,
        tone=(req.tone or "professional").lower(),
        company=req.company,
        role=req.role,
        email_type=(req.email_type or "application").lower(),
        recipient_name=req.recipient_name,
        job_link=req.job_link,
        date_applied=req.date_applied,
        availability=req.availability,
        extra_context=req.extra_context,
        drafts=max(1, min(int(req.drafts or 2), 4)),
    )

    return StreamingResponse(
        stream_ollama(prompt, temperature=0.6),
        media_type="text/plain; charset=utf-8",
    )


# ==============================
# AI INTERVIEWS — POST-INTERVIEW REPORT
# ==============================
# Called by resumeai-backend (app/utils/interviews.py::generate_report) once the
# live LiveKit interview has ended and the worker has posted the transcript.
# Input is text only (no audio ever reaches this service). Output is validated
# again by the backend before it is stored, so this endpoint focuses on getting
# a well-grounded JSON evaluation out of the model.

INTERVIEW_SCORE_WEIGHTS = {
    "relevance": 0.30,
    "evidence": 0.25,
    "structure": 0.15,
    "role_alignment": 0.20,
    "communication": 0.10,
}
INTERVIEW_TRANSCRIPT_MAX_CHARS = 24000
INTERVIEW_EVAL_VERSION = "interview-eval-v1"

_INTERVIEW_TYPE_LABELS = {
    "general": "general",
    "behavioural": "behavioural (STAR-style)",
    "technical": "technical",
    "hr_screening": "HR screening",
    "leadership": "leadership",
}
_SENIORITY_LABELS = {
    "entry": "entry level",
    "mid": "mid-level",
    "senior": "senior",
    "lead": "lead / manager",
}


class InterviewTurn(BaseModel):
    role: str
    text: str


class InterviewReportRequest(BaseModel):
    role_title: str
    interview_type: str = "general"
    seniority: str = "mid"
    duration_minutes: int = 15
    resume: Optional[Dict] = None
    job_description: Optional[str] = None
    transcript: List[InterviewTurn]


def _interview_resume_brief(resume: Optional[Dict]) -> str:
    """Compact, prompt-friendly text of the résumé snapshot (no contact details)."""
    if not isinstance(resume, dict) or not resume:
        return "No résumé was provided."
    lines: List[str] = []
    if resume.get("summary"):
        lines.append(f"Summary: {str(resume['summary'])[:600]}")
    for e in (resume.get("experiences") or [])[:8]:
        if not isinstance(e, dict):
            continue
        head = " — ".join([p for p in [e.get("role"), e.get("company")] if p])
        dates = " to ".join([p for p in [e.get("startDate"), e.get("endDate")] if p])
        lines.append(f"Experience: {head} ({dates})" if dates else f"Experience: {head}")
        for b in (e.get("bullets") or [])[:5]:
            lines.append(f"  - {str(b)[:220]}")
    for p in (resume.get("projects") or [])[:6]:
        if not isinstance(p, dict):
            continue
        lines.append(f"Project: {p.get('title') or 'Untitled project'}")
        for b in (p.get("bullets") or [])[:4]:
            lines.append(f"  - {str(b)[:220]}")
    skills: List[str] = []
    for cat in resume.get("skills") or []:
        if isinstance(cat, dict):
            skills.extend([str(x) for x in (cat.get("skills") or [])])
        elif isinstance(cat, str):
            skills.append(cat)
    if skills:
        lines.append("Skills: " + ", ".join(skills[:40]))
    for ed in (resume.get("education") or [])[:4]:
        if isinstance(ed, dict):
            lines.append("Education: " + " — ".join([p for p in [ed.get("degree"), ed.get("field"), ed.get("school")] if p]))
    return "\n".join(lines) if lines else "No résumé was provided."


def _interview_transcript_text(turns: List[InterviewTurn]) -> str:
    out: List[str] = []
    for i, t in enumerate(turns):
        role = "INTERVIEWER" if t.role == "assistant" else "CANDIDATE"
        text = " ".join(str(t.text or "").split())
        if text:
            out.append(f"[{i + 1}] {role}: {text}")
    text = "\n".join(out)
    if len(text) > INTERVIEW_TRANSCRIPT_MAX_CHARS:
        text = text[-INTERVIEW_TRANSCRIPT_MAX_CHARS:]
        cut = text.find("\n[")
        text = "(earlier turns truncated)\n" + (text[cut + 1:] if cut >= 0 else text)
    return text


def _build_interview_report_prompt(req: InterviewReportRequest) -> str:
    jd = " ".join(str(req.job_description or "").split())[:6000] or "Not provided."
    return f"""You are a fair, rigorous interview assessor. Evaluate the CANDIDATE in the spoken mock-interview transcript below and return ONLY a JSON object.

TARGET ROLE: {req.role_title}
INTERVIEW TYPE: {_INTERVIEW_TYPE_LABELS.get(req.interview_type, req.interview_type)}
SENIORITY BEING ASSESSED: {_SENIORITY_LABELS.get(req.seniority, req.seniority)}

CANDIDATE RÉSUMÉ (context only — do not score the résumé itself):
{_interview_resume_brief(req.resume)}

JOB DESCRIPTION:
{jd}

TRANSCRIPT (spoken, so expect fillers and small transcription errors):
{_interview_transcript_text(req.transcript)}

GROUNDING RULES (strict):
- Base every judgement ONLY on what the candidate actually said in the transcript. Never invent achievements, numbers, employers, or details that are not there. "evidence" must paraphrase or quote the candidate's own words.
- Score the CONTENT and STRUCTURE of the answers. Never score accent, grammar slips, fillers, transcription noise, personality, or anything about protected characteristics.
- Calibrate expectations to the seniority above (an entry-level candidate is not expected to talk about architecture strategy; a senior one is).
- Short or off-topic answers score low on relevance/evidence; do not pad them with charity.

WHAT TO RETURN:
1. "questions": the interviewer's substantive questions, in order (skip greetings, small talk, "are you ready", and the closing). Give each a stable id "q1", "q2", ... and a short category (e.g. "warm_up", "experience", "project", "technical", "behavioural", "role_fit", "follow_up"). Mark follow-ups with "is_follow_up": true.
2. "answers": one entry per question that the candidate actually answered. "transcript" = the candidate's answer, verbatim from the transcript (merge consecutive candidate turns; light cleanup of fillers is fine; do not paraphrase). Each answer has:
   - "scores": integers 0-100 for relevance (answers what was asked), evidence (specific, concrete, first-hand details and outcomes), structure (clear beginning/context/actions/result), role_alignment (how well it maps to the target role, JD and seniority), communication (clarity and concision of the spoken answer)
   - "evidence": 1-2 sentences pointing to what in the answer supports the scores
   - "worked": 1-3 short bullets of what was good
   - "improvements": 1-3 short, specific bullets
   - "improved_outline": 3-5 short bullets showing a stronger way to structure THIS answer using only facts the candidate mentioned
3. "report": overall view across all answers:
   - "scores": the same five dimensions, 0-100, reflecting the whole interview
   - "summary": 1-2 sentences of honest overall assessment
   - "strengths": 3 bullets, "improvements": 3 bullets (highest impact first), "action_plan": 4 concrete practice actions for before the real interview

Return valid JSON only, in exactly this shape:
{{
  "questions": [{{"id": "q1", "prompt": "...", "category": "warm_up", "is_follow_up": false}}],
  "answers": [{{"question_id": "q1", "transcript": "...", "scores": {{"relevance": 0, "evidence": 0, "structure": 0, "role_alignment": 0, "communication": 0}}, "evidence": "...", "worked": ["..."], "improvements": ["..."], "improved_outline": ["..."]}}],
  "report": {{"scores": {{"relevance": 0, "evidence": 0, "structure": 0, "role_alignment": 0, "communication": 0}}, "summary": "...", "strengths": ["..."], "improvements": ["..."], "action_plan": ["..."]}}
}}"""


def _clamp_interview_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def _normalize_interview_scores(scores) -> Dict[str, int]:
    scores = scores if isinstance(scores, dict) else {}
    return {k: _clamp_interview_score(scores.get(k, 0)) for k in INTERVIEW_SCORE_WEIGHTS}


def interview_overall_score(scores: Dict[str, int]) -> int:
    """Deterministic weighted overall — computed here in Python, never by the model."""
    return int(round(sum(scores.get(k, 0) * w for k, w in INTERVIEW_SCORE_WEIGHTS.items())))


@app.post("/interview/report")
def interview_report(req: InterviewReportRequest):
    candidate_turns = [t for t in req.transcript if t.role == "user" and str(t.text or "").strip()]
    if not candidate_turns:
        raise HTTPException(status_code=400, detail="Transcript contains no candidate answers.")

    prompt = _build_interview_report_prompt(req)
    try:
        raw = call_ollama(prompt)
        data = safe_json_parse(raw)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Interview evaluation failed: {e}")

    if not isinstance(data, dict):
        raise HTTPException(status_code=502, detail="Interview evaluation returned an invalid document.")

    answers = []
    for a in data.get("answers") or []:
        if not isinstance(a, dict):
            continue
        a["scores"] = _normalize_interview_scores(a.get("scores"))
        answers.append(a)
    report = data.get("report") if isinstance(data.get("report"), dict) else {}
    if not isinstance(report.get("scores"), dict) and answers:
        report["scores"] = {
            k: round(sum(a["scores"][k] for a in answers) / len(answers)) for k in INTERVIEW_SCORE_WEIGHTS
        }
    report["scores"] = _normalize_interview_scores(report.get("scores"))
    report["overall_score"] = interview_overall_score(report["scores"])

    return {
        "questions": data.get("questions") or [],
        "answers": answers,
        "report": report,
        "evaluation_version": INTERVIEW_EVAL_VERSION,
    }
